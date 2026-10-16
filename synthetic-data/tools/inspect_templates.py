from docx import Document
from pathlib import Path
from zipfile import ZipFile
from lxml import etree

for path in sorted(Path("data/templates").glob("template*.docx")):
    print(f"\n=== {path.name} ===")
    doc = Document(path)
    sec = doc.sections[0]
    print("page", sec.page_width, sec.page_height, "margins", sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin)
    print("paragraphs", len(doc.paragraphs), "tables", len(doc.tables), "shapes", len(doc.inline_shapes))
    for p in doc.paragraphs:
        if p.text.strip():
            print("P:", p.text)
    for i, t in enumerate(doc.tables):
        print("TABLE", i)
        for row in t.rows:
            print(" | ".join(c.text.replace("\n", " / ") for c in row.cells))
    with ZipFile(path) as z:
        xml = etree.fromstring(z.read("word/document.xml"))
        texts = xml.xpath("//*[local-name()='t']/text()")
        print("xml_text:", " || ".join(texts))
        print("document parts:", len(z.namelist()))
