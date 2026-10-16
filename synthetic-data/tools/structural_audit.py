from pathlib import Path
from zipfile import ZipFile
from docx import Document
from lxml import etree

for f in [Path("data/resume_09_han_zimo_ios_template04.docx"), Path("data/resume_10_lin_ruoxi_customer_success_template03.docx")]:
    doc = Document(f)
    sec = doc.sections[0]
    with ZipFile(f) as z:
        root = etree.fromstring(z.read("word/document.xml"))
        media = [n for n in z.namelist() if n.startswith("word/media/") and n.endswith(".png")]
        text = "".join(root.xpath("//*[local-name()='t']/text()"))
    print(f.name)
    print("  page", sec.page_width, sec.page_height, "margins", sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin)
    print("  paragraphs", len(doc.paragraphs), "tables", len(doc.tables), "media", len(media), "characters", len(text), "size", f.stat().st_size)
    print("  old_content_removed", not any(x in text for x in ["某某某", "医护岗位", "上海赫尔森康复医院", "XXX@ibaotu.com"]))
