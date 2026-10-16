from docx import Document
from pathlib import Path
import zipfile

data = Path("data")
files = sorted(data.glob("resume_*.docx"))
print("count", len(files))
for f in files:
    d = Document(f)
    txt = "\n".join(p.text for p in d.paragraphs)
    txt += "\n" + "\n".join(c.text for t in d.tables for row in t.rows for c in row.cells)
    with zipfile.ZipFile(f) as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]
    footer = "\n".join(p.text for p in d.sections[0].footer.paragraphs)
    print(f"{f.name} | paras={len(d.paragraphs)} tables={len(d.tables)} images={len(media)} chars={len(txt)} footer={('模拟简历' in footer)}")
