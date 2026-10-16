from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from docx import Document

files = [
    Path("data/resume_09_han_zimo_ios_template04.docx"),
    Path("data/resume_10_lin_ruoxi_customer_success_template03.docx"),
]
old_markers = ["免费简历", "某某某", "医护岗位", "上海赫尔森康复医院", "IOS工程师", "XXX@ibaotu.com"]
for path in files:
    doc = Document(path)
    with ZipFile(path) as z:
        root = etree.fromstring(z.read("word/document.xml"))
        all_text = "".join(root.xpath("//*[local-name()='t']/text()"))
        media = [n for n in z.namelist() if n.startswith("word/media/") and n.lower().endswith(".png")]
    print(path.name, "paragraphs", len(doc.paragraphs), "tables", len(doc.tables), "media", len(media), "bytes", path.stat().st_size)
    for marker in old_markers:
        if marker in all_text:
            print("OLD_MARKER_PRESENT", marker)
    for marker in ["韩子墨", "iOS 开发工程师", "林若溪", "客户成功经理"]:
        if marker in all_text:
            print("FOUND", marker)
