import base64
import io
import json
import zipfile
from pathlib import Path

import pytest
import pymupdf
from docx import Document
from fastapi.testclient import TestClient
from app.api import create_app
from app.config import ROOT, Settings
from app.service import Service
from app.retrieval import search
from app.parsers import validate_file

@pytest.fixture
def service(tmp_path):
    return Service(Settings(data_dir=tmp_path,soffice_path='',embedding_provider='local-hash',embedding_model='local-hash-v1',chat_model=''))

def resume(name='张三',skill='Java Spring Boot MySQL',email='zhang@example.com'):
    doc=Document()
    for text in [name,'后端开发工程师','现居：上海','5年工作经验',f'邮箱：{email}','个人概述',f'从事后端开发，熟悉 {skill}，负责企业服务平台研发与需求分析。','教育背景','2015.09-2019.06 某某大学 计算机科学 本科','工作经历','2019.07-至今 上海示例科技有限公司 后端工程师',f'负责 {skill} 平台的需求分析、数据库设计、接口开发和上线维护，优化查询性能。']:
        doc.add_paragraph(text)
    stream=io.BytesIO(); doc.save(stream)
    return stream.getvalue()

def ingest(service,data,filename='resume.docx'):
    return service.process(service.register(filename,data))

def test_duplicate_version_and_active_indexes(service):
    first=ingest(service,resume())
    assert first['status']=='partial_success'
    duplicate=ingest(service,resume())
    assert duplicate['status']=='rejected_duplicate'
    assert duplicate['version_no']==0
    changed=ingest(service,resume(skill='Rust Tokio'))
    assert changed['candidate_id']==first['candidate_id']
    assert changed['version_no']==2
    assert service.document(first['document_id'])['is_active']==0
    with service.store.connect() as db:
        assert db.execute('SELECT count(*) FROM source_documents WHERE is_active=1').fetchone()[0]==1
        assert db.execute('SELECT count(*) FROM chunks_fts WHERE text LIKE ?',('%java%',)).fetchone()[0]==0
        assert db.execute('SELECT count(*) FROM chunk_vectors').fetchone()[0]==db.execute('SELECT count(*) FROM chunks WHERE is_active=1').fetchone()[0]

def test_failed_embedding_preserves_previous(service,monkeypatch):
    first=ingest(service,resume())
    def fail(_): raise ValueError('Embedding 服务失败')
    monkeypatch.setattr(service.models,'embed',fail)
    second=ingest(service,resume(skill='Rust Tokio'))
    assert second['status']=='failed'
    assert service.candidates()[0]['latest_document_id']==first['document_id']

def test_privacy_filters_and_citations(service):
    first=ingest(service,resume())
    answer=search(service,'上海 Java 5 年以上')
    assert answer['results'][0]['name']=='张三'
    assert answer['results'][0]['evidence']
    assert 'zhang@example.com' not in json.dumps(answer)
    assert not search(service,'北京 Java 5 年以上')['results']
    assert not search(service,'上海 Java 20 年以上')['results']
    assert not search(service,'只要男性')['results']
    assert not search(service,'量子芯片光刻专家')['results']
    archive=service.archive(first['document_id'])
    assert archive['profile']['email']=='[邮箱已隐藏]'

def test_file_validation_and_scanned_pdf(service,tmp_path):
    bad=ingest(service,b'not a pdf','fake.pdf')
    assert bad['status']=='failed'
    buffer=io.BytesIO()
    with zipfile.ZipFile(buffer,'w') as z:
        z.writestr('../attack','x')
    bad=ingest(service,buffer.getvalue(),'evil.docx')
    assert bad['status']=='failed'
    assert '不安全路径' in bad['error']
    pdf=pymupdf.open(); pdf.new_page()
    empty=ingest(service,pdf.tobytes(),'scanned.pdf')
    assert empty['status']=='failed' and 'OCR' in empty['error']
    assert service.candidates()==[]

def test_sensitive_table_not_in_knowledge(service):
    doc=Document(io.BytesIO(resume()))
    doc.add_paragraph('紧急联系人')
    table=doc.add_table(rows=2,cols=2)
    table.cell(0,0).text='姓名'; table.cell(0,1).text='电话'
    table.cell(1,0).text='王保密'; table.cell(1,1).text='13912345678'
    buf=io.BytesIO();doc.save(buf)
    result=ingest(service,buf.getvalue())
    rows=service.store.rows('SELECT clean_text FROM knowledge_documents')
    assert rows and '王保密' not in rows[0]['clean_text']
    assert service.archive(result['document_id'])['tables']

def test_baseline_all_formats(service):
    results=[]
    for path in sorted((ROOT/'synthetic-data/cv').glob('*')):
        result=ingest(service,path.read_bytes(),path.name)
        assert result['status'] in {'succeeded','partial_success','rejected_duplicate'},result
        results.append(result)
    assert len(results)==12
    pdf03=next(r for r in results if r['filename']=='resume_03_zhao_xinyi_hr_admin.pdf')
    assert pdf03['status']=='rejected_duplicate'
    assert len(service.candidates())==10
    names={r['name'] for r in service.candidates()}
    assert {'林晨','何婉宁','林若溪','韩子墨','郑凯'} <= names
    for query,expected in [('熟悉 Java、Spring Boot、MySQL 的后端应届生','林晨'),('带领测试团队并做过接口自动化','郑凯'),('上海 5 年以上 SaaS 产品经验，做过客户服务或知识库','何婉宁')]:
        result=search(service,query)
        assert result['results'][0]['name']==expected
        assert len({r['candidate_id'] for r in result['results']})==len(result['results'])
    for r in results:
        if r['status']=='rejected_duplicate': continue
        archive=service.archive(r['document_id'])
        assert archive['assets']['portrait'],r['filename']
        for portrait in archive['assets']['portrait']:
            assert base64.b64decode(Path(portrait['base64_path']).read_text())==Path(portrait['path']).read_bytes()
        if r['filename'].endswith('.pdf'):
            assert archive['assets']['snapshots']
            assert Path(archive['assets']['snapshots'][0]['path']).is_file()

def test_api(tmp_path):
    app=create_app(Settings(data_dir=tmp_path),mount_ui=False)
    with TestClient(app) as client:
        assert client.get('/api/health').status_code==200
        assert client.get('/api/documents/missing').status_code==404
        assert client.post('/api/search',json={'query':'','top_k':50}).status_code==422
        r=client.post('/api/documents',files=[('files',('r.docx',resume()))])
        assert r.status_code==202
    # Lifespan waits for submitted ingestion on shutdown.
    assert len(app.state.service.candidates())==1
