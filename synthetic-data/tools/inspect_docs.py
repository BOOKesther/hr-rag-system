from docx import Document
import glob
import os

for f in glob.glob('data/*.docx'):
    print(f"\n=== {os.path.basename(f)} ===")
    d = Document(f)
    for p in d.paragraphs:
        if p.text.strip():
            print('P:', p.text)
    for ti, t in enumerate(d.tables):
        print('TABLE', ti)
        for row in t.rows:
            print(' | '.join(c.text.replace('\n', ' / ') for c in row.cells))
    print('inline_shapes', len(d.inline_shapes))
